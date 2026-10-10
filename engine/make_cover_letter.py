#!/usr/bin/env python3
"""Write a cover letter .docx (and PDF via LibreOffice) from a JSON spec.

Usage: python3 make_cover_letter.py spec.json
spec: {"filename": "Your_Name_Company_Cover_Letter.docx", "date": "September 13, 2026",
       "recipient_lines": ["Hiring Team", "Company"], "greeting": "Dear Hiring Team,",
       "paragraphs": ["...", "..."], "closing": "Sincerely,", "signature": "Your Name"}
The header (name, contact line, LinkedIn and website) comes from the profile. filename defaults
to <file_stem>_Cover_Letter.docx and signature to the profile's cover_letter.signature.
Format: US Letter, 1 inch margins, the profile's resume font at 11pt, single spaced, 8pt after paragraphs.
"""
import json
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

_here = os.path.dirname(os.path.abspath(__file__))


def _find_server():
    """server/userprofile.py, found by walking up: this script runs from <repo>/engine or from a
    copy in <data>/applications/<job>/_build/."""
    d = _here
    for _ in range(8):
        if os.path.isfile(os.path.join(d, "server", "userprofile.py")):
            return os.path.join(d, "server")
        if os.path.dirname(d) == d:
            break
        d = os.path.dirname(d)
    sys.exit("server/userprofile.py not found above " + _here)


try:
    import userprofile as up  # the server passes PYTHONPATH, which covers a data folder outside the repo
except ImportError:
    sys.path.insert(0, _find_server())
    import userprofile as up  # noqa: E402


def _short(url):
    """https://www.linkedin.com/in/x/ -> linkedin.com/in/x, the way it reads on paper."""
    return re.sub(r"^https?://(www\.)?", "", (url or "").strip()).rstrip("/")


def header_lines():
    ident = up.get()["identity"]
    contact = [up.contact_line()] + [_short(ident.get(k)) for k in ("linkedin", "website")]
    return ident.get("name") or "", " | ".join(x for x in contact if x)


def main(spec_path):
    spec = json.load(open(spec_path))
    # Same rule as build.py: finished letters go to the application folder, not into _build/.
    out_dir = os.path.dirname(_here) if os.path.basename(_here) == "_build" else os.path.dirname(os.path.abspath(spec_path))
    font = (up.get().get("resume") or {}).get("font") or "Times New Roman"
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    style = doc.styles["Normal"]
    style.font.name = font
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", font)
    doc.core_properties.author = up.get()["identity"].get("name") or ""
    doc.core_properties.last_modified_by = ""

    def para(text, after=8, bold=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.bold = bold
        return p

    name, contact = header_lines()
    signature = spec.get("signature") or up.get()["cover_letter"].get("signature") or name
    para(name, after=0, bold=True)
    para(contact, after=12)
    para(spec.get("date", ""), after=12)
    for line in spec.get("recipient_lines", []):
        para(line, after=0)
    para("", after=4)
    para(spec.get("greeting", "Dear Hiring Team,"), after=8)
    for ptxt in spec["paragraphs"]:
        para(ptxt.replace("—", ","), after=8)
    para(spec.get("closing", "Sincerely,"), after=18)
    para(signature, after=0)

    docx_path = os.path.join(out_dir, spec.get("filename") or up.file_stem() + "_Cover_Letter.docx")
    doc.save(docx_path)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path],
                   capture_output=True, text=True, timeout=180)
    pdf = docx_path[:-5] + ".pdf"
    pages = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout
    print(json.dumps({"docx": docx_path, "pdf": pdf,
                      "pages": next((int(l.split()[-1]) for l in pages.splitlines() if l.startswith("Pages:")), None)}))


if __name__ == "__main__":
    main(sys.argv[1])
